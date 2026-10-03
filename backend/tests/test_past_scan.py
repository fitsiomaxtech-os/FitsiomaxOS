"""Auto Scan (past_scan.py), and the custom sheet it opens up (past_custom.py): every tab and
header of an upload listed, each tab and column turned on or off before anything is read, and
a list of people that is none of the three known layouts read column by column.

Synthetic people only -- no row of a real sheet goes into a file.
"""
import os
import sys
from datetime import date, datetime

import pytest

# Same note as the files beside this one: pytest with no __init__.py puts this directory on
# the path rather than the backend root.
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import past_custom  # noqa: E402
from past_data import PastDataError, read_workbook, summary  # noqa: E402
from past_os import COLUMNS  # noqa: E402
from past_scan import header_row, scan_file  # noqa: E402
from test_past_data import write_register  # noqa: E402
from test_past_os import _workbook  # noqa: E402

# The example the feature was asked for with: a "Data" workbook, four tabs, Sheet1 with these.
HEADERS = ["Name", "Phone", "Email", "City", "state", "country", "DoB", "Company name", "work type",
           "Active or Inactive"]


def data_file(path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    s1 = wb.create_sheet("Sheet1")
    s1.append(HEADERS)
    s1.append(["Test Arun", 9000000001, "arun@example.com", "Chennai", "Tamil Nadu", "India",
               datetime(1990, 5, 20), "Acme", "IT", "Active"])
    s1.append(["Test Priya", "+91 90000 00002", "priya@example.com", "Madurai", "Tamil Nadu", "India",
               None, "Beta", "Sales", "Inactive"])
    s1.append([None] * 10)
    s1.append([None, None, None, "Coimbatore"])
    s2 = wb.create_sheet("Sheet2")
    s2.append(["Contacts exported 30 Sep"])
    s2.append([])
    s2.append(["Client Name", "Contact 1", "Remarks"])
    s2.append(["Test Arun", "9000000001", "Came back in May"])
    s2.append(["Test Kavi", "9000000003", ""])
    wb.create_sheet("Sheet3")
    s4 = wb.create_sheet("Sheet4")
    s4.sheet_state = "hidden"
    s4.append(["Name", "Phone"])
    s4.append(["Test Hidden", "9000000004"])
    wb.save(path)
    return path


def tabs_of(scanned):
    return {t["name"]: t for t in scanned["tabs"]}


def columns_of(tab):
    return {c["header"]: c for c in tab["columns"]}


def picks(scanned, off_tabs=(), off=(), fields=None):
    """The scan's own defaults as the screen sends them, less what a test turns off."""
    fields = fields or {}
    out = []
    for tab in scanned["tabs"]:
        if not tab["on"] or tab["name"] in off_tabs:
            continue
        out.append({"tab": tab["name"], "columns": [
            {"index": c["index"], "field": fields.get((tab["name"], c["header"]), c["field"])}
            for c in tab["columns"] if c["on"] and (tab["name"], c["header"]) not in off
        ]})
    return out


# ------------------------------------------------------------------------------ the scan

def test_scan_lists_every_tab_and_header(tmp_path):
    scanned = scan_file(data_file(tmp_path / "Data.xlsx"))
    assert scanned["layout"] == "custom"
    tabs = tabs_of(scanned)
    assert list(tabs) == ["Sheet1", "Sheet2", "Sheet3", "Sheet4"]

    sheet1 = tabs["Sheet1"]
    assert (sheet1["header_row"], sheet1["rows"], sheet1["on"]) == (1, 3, True)
    assert [c["header"] for c in sheet1["columns"]] == HEADERS
    assert all(c["on"] and c["used"] and not c["required"] for c in sheet1["columns"])
    cols = columns_of(sheet1)
    assert cols["Phone"]["samples"] == ["9000000001", "+91 90000 00002"]
    assert cols["DoB"]["samples"] == ["1990-05-20"] and cols["DoB"]["filled"] == 1
    assert cols["City"]["filled"] == 3


def test_scan_suggests_an_os_field_off_each_header(tmp_path):
    tabs = tabs_of(scan_file(data_file(tmp_path / "Data.xlsx")))
    fields = {h: c["field"] for h, c in columns_of(tabs["Sheet1"]).items()}
    assert fields == {
        "Name": "name", "Phone": "phone", "Email": "email", "City": "city", "state": "state",
        "country": "", "DoB": "dob", "work type": "", "Active or Inactive": "status",
        # "name" is in it, and it is not the client's name.
        "Company name": "",
    }
    # Below a title, and a column of numbers under a header the OS does not know.
    sheet2 = tabs["Sheet2"]
    assert sheet2["header_row"] == 3 and sheet2["rows"] == 2
    assert {h: c["field"] for h, c in columns_of(sheet2).items()} == {
        "Client Name": "name", "Contact 1": "phone", "Remarks": "notes"}


def test_scan_leaves_an_empty_tab_and_a_hidden_one_off(tmp_path):
    tabs = tabs_of(scan_file(data_file(tmp_path / "Data.xlsx")))
    assert (tabs["Sheet3"]["used"], tabs["Sheet3"]["on"], tabs["Sheet3"]["note"]) == (False, False, "Empty")
    assert (tabs["Sheet4"]["hidden"], tabs["Sheet4"]["used"], tabs["Sheet4"]["on"]) == (True, True, False)


def test_header_row_is_the_headers_not_a_row_of_words_under_them():
    rows = [
        ("Camp list",),
        ("Name", "City", "Occupation", "Phone"),
        ("Test Arun", "Chennai", "Engineer", "9000000001"),
        ("Test Priya", "Madurai", "Teacher", "Salem"),
    ]
    assert header_row(rows) == 1
    assert header_row([(1, 2, 3)]) is None


# ------------------------------------------------------------------ fetching a custom sheet

def test_fetch_reads_only_the_columns_left_on(tmp_path):
    path = data_file(tmp_path / "Data.xlsx")
    scanned = scan_file(path)
    data = read_workbook(path, picks(scanned, off={("Sheet1", "work type")}))
    assert data.layout == "custom" and data.types == ["lead"]
    assert data.tabs == ["Sheet1", "Sheet2"] and data.tab_rows == {"Sheet1": 3, "Sheet2": 2}
    by_name = {c["name"]: c for c in data.clients}
    assert sorted(by_name) == ["Test Arun", "Test Kavi", "Test Priya"]

    arun = by_name["Test Arun"]
    assert (arun["phone_normalized"], arun["email"], arun["city"], arun["state"]) == (
        "9000000001", "arun@example.com", "Chennai", "Tamil Nadu")
    assert arun["dob"] == "1990-05-20" and arun["age"] == past_custom._age("1990-05-20", date.today())
    assert arun["latest_status"] == "Active" and arun["excel_id"] == "Sheet1 · row 2"
    # Kept as written, under its own header; the column turned off is nowhere.
    assert arun["extra"] == {"country": "India", "Company name": "Acme"}
    # The same person on Sheet2 is the same client, filled in from there.
    assert arun["notes"] == "Came back in May"
    assert by_name["Test Priya"]["phone_normalized"] == "9000000002"

    codes = {f.code: f for f in data.findings}
    assert set(codes) == {"no_name_row", "merged_rows"}
    assert codes["no_name_row"].excel_id == "Sheet1 · row 5"
    assert codes["merged_rows"].detail == "Test Arun: also Sheet2 · row 4"
    assert data.columns[0]["tab"] == "Sheet1" and len(data.columns[0]["columns"]) == 9
    assert {"header": "Company name", "field": ""} in data.columns[0]["columns"]
    assert summary(data)["layout"] == "custom"


def test_a_column_can_be_given_another_field(tmp_path):
    path = data_file(tmp_path / "Data.xlsx")
    scanned = scan_file(path)
    data = read_workbook(path, picks(scanned, fields={("Sheet1", "work type"): "occupation",
                                                      ("Sheet1", "City"): ""}))
    arun = next(c for c in data.clients if c["name"] == "Test Arun")
    assert arun["occupation"] == "IT" and arun["city"] == ""
    assert arun["extra"]["City"] == "Chennai"


def test_a_phone_is_checked_only_where_the_tab_has_a_phone_column():
    data = past_custom.build([
        ("Emails", [{"fields": {"name": "Test Arun", "email": "a@example.com"}, "extra": {}, "_row": 2}]),
        ("Calls", [{"fields": {"name": "Test Priya", "phone": "98765"}, "extra": {}, "_row": 2}]),
    ])
    flags = {c["name"]: c["flags"] for c in data.clients}
    assert flags == {"Test Arun": [], "Test Priya": ["check_phone"]}


def test_one_number_under_two_names_is_two_people():
    data = past_custom.build([("Sheet1", [
        {"fields": {"name": "Test Arun", "phone": "9000000001"}, "extra": {}, "_row": 2},
        {"fields": {"name": "Test Meena", "phone": "9000000001"}, "extra": {}, "_row": 3},
    ])])
    assert [c["name"] for c in data.clients] == ["Test Arun", "Test Meena"]
    assert all("shared_phone" in c["flags"] for c in data.clients)


def test_a_current_stage_column_names_the_os_stage():
    data = past_custom.build([("Sheet1", [
        {"fields": {"name": "Test Arun", "current_stage": "follow up"}, "extra": {}, "_row": 2},
        {"fields": {"name": "Test Priya", "current_stage": "Maybe later"}, "extra": {}, "_row": 3},
    ])])
    arun, priya = data.clients
    assert arun["current_stage"] == "Follow Up"
    assert priya["current_stage"] == "" and "unknown_stage" in priya["flags"]


@pytest.mark.parametrize("change, message", [
    ({"off_tabs": ("Sheet1", "Sheet2")}, "every tab is turned off"),
    ({"off": {("Sheet2", "Client Name"), ("Sheet2", "Contact 1")}}, "give one column Name or Phone"),
    ({"fields": {("Sheet1", "City"): "name"}}, "Name is given to two columns"),
])
def test_fetch_refuses_picks_that_cannot_be_read(tmp_path, change, message):
    path = data_file(tmp_path / "Data.xlsx")
    with pytest.raises(PastDataError, match=message):
        read_workbook(path, picks(scan_file(path), **change))


def test_fetch_refuses_a_tab_or_a_column_the_scan_never_listed(tmp_path):
    path = data_file(tmp_path / "Data.xlsx")
    with pytest.raises(PastDataError, match="no 'Nope' tab"):
        read_workbook(path, [{"tab": "Nope", "columns": []}])
    with pytest.raises(PastDataError, match="no table to read"):
        read_workbook(path, [{"tab": "Sheet3", "columns": []}])
    with pytest.raises(PastDataError, match="has no column 40"):
        read_workbook(path, [{"tab": "Sheet1", "columns": [{"index": 40, "field": "name"}]}])
    with pytest.raises(PastDataError, match="not an OS field"):
        read_workbook(path, [{"tab": "Sheet1", "columns": [{"index": 0, "field": "salary"}]}])
    with pytest.raises(PastDataError, match="could not be read"):
        read_workbook(path, {"Sheet1": []})


def test_the_terminal_reads_a_custom_sheet_as_the_scan_first_shows_it(tmp_path):
    data = read_workbook(data_file(tmp_path / "Data.xlsx"))
    assert data.layout == "custom" and data.columns == []
    # Sheet4 is hidden, and the scan leaves it off.
    assert sorted(c["name"] for c in data.clients) == ["Test Arun", "Test Kavi", "Test Priya"]
    arun = next(c for c in data.clients if c["name"] == "Test Arun")
    assert arun["extra"] == {"country": "India", "Company name": "Acme", "work type": "IT"}


# ------------------------------------------------------------------- the three known layouts

def test_the_register_marks_what_it_cannot_do_without(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "register.xlsx"
    write_register(path)
    wb = load_workbook(path)
    wb.create_sheet("Controls").append(["Service_Interest", "Offline Group Fitness"])
    wb.save(path)

    scanned = scan_file(path)
    assert scanned["layout"] == "register" and scanned["fields"] == []
    tabs = tabs_of(scanned)
    assert all(tabs[t]["required"] and tabs[t]["on"] for t in
               ("Patient Master", "Enrollment Register", "Payment Tracker", "Prospect Pipeline"))
    assert (tabs["Controls"]["used"], tabs["Controls"]["note"]) == (False, "Not part of the register")
    master = columns_of(tabs["Patient Master"])
    assert tabs["Patient Master"]["header_row"] == 2
    assert {h for h, c in master.items() if c["required"]} == {"Patient ID", "Full Name", "Phone"}
    assert master["Gender"]["on"] and not master["Gender"]["required"]

    whole = read_workbook(path)
    data = read_workbook(path, picks(scanned, off={("Patient Master", "Gender")}))
    assert [c["gender"] for c in whole.clients] == ["Female", "Female"]
    assert [c["gender"] for c in data.clients] == ["", ""]
    assert len(data.payments) == len(whole.payments) == 1

    with pytest.raises(PastDataError, match="Phone is needed"):
        read_workbook(path, picks(scanned, off={("Patient Master", "Phone")}))
    with pytest.raises(PastDataError, match="'Prospect Pipeline' is needed to read the register"):
        read_workbook(path, picks(scanned, off_tabs={"Prospect Pipeline"}))


def test_the_os_data_workbook_can_leave_a_tab_out(tmp_path):
    path = _workbook(tmp_path / "OSDATAX.xlsx", {
        "Leads": (COLUMNS["Leads"], [{"excel_id": "FM-1", "name": "Test Arun", "phone": "9000000001",
                                     "current_stage": "RNR", "city": "Chennai"}]),
        "Payments": (COLUMNS["Payments"], [{"client_excel_id": "FM-1", "status": "Paid", "amount": 700,
                                           "paid_date": datetime(2026, 9, 1), "mode": "Cash"}]),
    })
    scanned = scan_file(path)
    assert scanned["layout"] == "os"
    leads = columns_of(tabs_of(scanned)["Leads"])
    assert {h for h, c in leads.items() if c["required"]} == {"Patient ID", "Name", "Phone", "Current Stage"}

    data = read_workbook(path, picks(scanned, off_tabs={"Payments"}, off={("Leads", "City")}))
    assert data.tabs_missing[-1] == "Payments" and data.payments == [] and data.types == ["lead"]
    assert data.clients[0]["city"] == ""
    with pytest.raises(PastDataError, match="Current Stage is needed"):
        read_workbook(path, picks(scanned, off={("Leads", "Current Stage")}))


def test_the_os_data_instalment_columns_are_read_and_can_be_turned_off(tmp_path):
    columns = {k: COLUMNS["Payments"][k] for k in ("client_excel_id", "name", "consultation_fee", "consultation_date")}
    for n in (1, 2):
        columns.update({f"instalment_{n}_amount": f"Instalment {n} Amount",
                        f"instalment_{n}_paid_date": f"Instalment {n} Paid Date"})
    path = _workbook(tmp_path / "OSDATAX.xlsx", {
        "Leads": (COLUMNS["Leads"], [{"excel_id": "FM-1", "name": "Test Arun", "phone": "9000000001",
                                     "current_stage": "RNR"}]),
        "Payments": (columns, [{"client_excel_id": "FM-1", "consultation_fee": 500,
                                "consultation_date": datetime(2026, 9, 1), "instalment_1_amount": 3000,
                                "instalment_1_paid_date": datetime(2026, 9, 2), "instalment_2_amount": 3000,
                                "instalment_2_paid_date": datetime(2026, 9, 9)}]),
    })
    scanned = scan_file(path)
    payments = columns_of(tabs_of(scanned)["Payments"])
    assert all(c["used"] and c["on"] for c in payments.values())
    assert {h for h, c in payments.items() if c["required"]} == {"Patient ID"}

    assert len(read_workbook(path, picks(scanned)).payments) == 3
    data = read_workbook(path, picks(scanned, off={("Payments", "Instalment 2 Amount"),
                                                   ("Payments", "Instalment 2 Paid Date")}))
    assert [p["excel_id"] for p in data.payments] == ["Payments · row 2 · Consultation", "Payments · row 2 · Instalment 1"]


def test_a_revenue_sheet_reads_without_a_column_turned_off(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "May 26"
    ws.append(["S.NO", "DATE", "NAME", "MOBILE NUMBER", "LOCATION", "DESCRIPTION", "PAYMENT", "MODE"])
    ws.append([1, datetime(2026, 5, 2), "Test Arun", "9000000001", "Anna Nagar", "Cons", 500, "Cash"])
    path = tmp_path / "revenue.xlsx"
    wb.save(path)

    scanned = scan_file(path)
    assert scanned["layout"] == "revenue"
    cols = columns_of(tabs_of(scanned)["May 26"])
    assert {h for h, c in cols.items() if c["required"]} == {"NAME", "MOBILE NUMBER", "PAYMENT"}
    assert not cols["S.NO"]["used"] and not cols["S.NO"]["on"]

    assert read_workbook(path).clients[0]["notes"] == "Location: Anna Nagar"
    data = read_workbook(path, picks(scanned, off={("May 26", "LOCATION")}))
    assert data.clients[0]["notes"] == "" and data.payments[0]["amount_paid"] == 500
