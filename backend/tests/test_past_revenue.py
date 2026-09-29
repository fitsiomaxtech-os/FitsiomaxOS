"""The revenue-sheet reader (past_revenue.py): a branch's month tabs, one payment a row.

Synthetic people only -- no row of a real sheet goes into a file.
"""
from datetime import datetime

import pytest

import past_data
from past_data import PastDataError, read_workbook
from past_revenue import _one_person, build


def row(n, name, phone, paid, date=datetime(2026, 5, 2), **extra):
    record = {"name": name, "phone": phone, "payment": paid, "date": date, "mode": "Gpay", "next": [], "_row": n}
    record.update(extra)
    return record


def by_name(clients):
    return {c["name"]: c for c in clients}


def codes(data):
    return sorted({f.code for f in data.findings})


def test_one_mobile_is_one_client_across_months():
    data = build([
        ("May 26", [row(2, "Test Alpha", "9000000001", 600, description="Cons")]),
        ("June 26", [row(5, "Test Alpha", 9000000001, 11200, date=datetime(2026, 6, 3), description="Physio session")]),
    ])
    assert len(data.clients) == 1
    client = data.clients[0]
    assert client["excel_id"] == "May 26 · row 2"
    assert client["treatments_count"] == 2 and client["paid_total"] == 11800
    assert (client["registration_date"], client["last_treatment_date"]) == ("2026-05-02", "2026-06-03")
    assert {t["tab"] for t in data.treatments} == {"May 26", "June 26"}
    assert all(p["client_id"] == client["id"] for p in data.payments)


@pytest.mark.parametrize("first,second", [
    ("Mohammed Hakeem", "Hakeem"),       # the surname alone
    ("Iqbal", "Iqibal"),                 # a letter too many
    ("Syed Ibrahim", "SYED IBHRAHIM"),   # a typo, in capitals
    ("Anita", "Anita Arun"),             # the register's own rule
])
def test_one_person_written_two_ways_on_one_mobile(first, second):
    assert _one_person(first, second)
    data = build([("May 26", [row(2, first, "9000000002", 500), row(3, second, "9000000002", 500)])])
    assert len(data.clients) == 1
    assert "shared_phone" not in codes(data)


@pytest.mark.parametrize("first,second", [
    ("Shaik Hameed", "Shaik Ahmed"),     # relatives share most of a name
    ("Test Mother", "Test Daughter"),
    ("Priya", "Divya"),
])
def test_two_people_on_one_mobile_stay_two_and_are_flagged(first, second):
    assert not _one_person(first, second)
    data = build([("May 26", [row(2, first, "9000000003", 500), row(3, second, "9000000003", 500)])])
    assert len(data.clients) == 2
    assert all("shared_phone" in c["flags"] for c in data.clients)
    a, b = data.clients
    assert a["shared_phone_with"] == [b["excel_id"]] and b["shared_phone_with"] == [a["excel_id"]]


def test_a_row_without_a_usable_mobile_is_its_own_client():
    data = build([("May 26", [row(2, "Test Beta", None, 500), row(3, "Test Beta", "", 700), row(4, "Test Beta", "98400", 900)])])
    assert len(data.clients) == 3
    assert all("check_phone" in c["flags"] for c in data.clients)


def test_a_nameless_row_on_a_known_mobile_joins_that_client():
    data = build([("May 26", [row(2, "", "9000000004", 500), row(3, "Test Gamma", "9000000004", 700)])])
    assert len(data.clients) == 1
    assert data.clients[0]["name"] == "Test Gamma"


def test_the_months_total_line_is_not_a_payment_and_is_reported():
    data = build([("May 26", [
        row(2, "Test Delta", "9000000005", 1000),
        {"name": None, "phone": None, "payment": 1000, "date": None, "next": [], "_row": 60},
        {"name": None, "phone": None, "payment": None, "date": None, "next": [], "_row": 61},  # empty: nothing
    ])])
    assert len(data.payments) == 1
    left_out = [f for f in data.findings if f.code == "no_name_row"]
    assert [f.excel_id for f in left_out] == ["May 26 · row 60"]
    assert "no_name_row" in past_data.NOT_IMPORTED


def test_balance_and_next_payment_are_kept_as_written_and_nothing_is_owed():
    data = build([("May 26", [row(
        2, "Test Epsilon", "9000000006", 11200, balance=11200, remarks="2w(4th week)", slot="5-6PM",
        next=[("NEXT PYMT DATE-MAY", "16-05-2026-paid"), ("NEXT PYMT DATE-JUNE", None)],
    )])])
    treatment, payment = data.treatments[0], data.payments[0]
    assert treatment["notes"] == "2w(4th week) · Balance: 11200 · NEXT PYMT DATE-MAY: 16-05-2026-paid · Slot 5-6PM"
    assert payment["balance"] == 11200
    assert (payment["state"], payment["outstanding"]) == ("paid", 0)
    assert data.clients[0]["outstanding_total"] == 0


def test_a_zero_balance_is_not_a_note():
    data = build([("May 26", [row(2, "Test Zeta", "9000000007", 800, balance=0)])])
    assert data.treatments[0]["notes"] == ""


@pytest.mark.parametrize("written,mode,flagged", [
    ("Gpay", "upi", False),
    ("G PAY", "upi", False),
    ("CCARD", "card", False),
    ("C CARD", "card", False),
    ("debit card", "card", False),
    ("AC transfer", "account_transfer", False),
    ("Cash & Gpay", "", False),   # paid two ways: kept as written, not filed under one
    ("Cash and Gpay", "", False),
    ("", "", True),
])
def test_payment_modes(written, mode, flagged):
    data = build([("May 26", [row(2, "Test Eta", "9000000008", 800, mode=written)])])
    payment = data.payments[0]
    assert payment["mode"] == mode and payment["mode_as_written"] == written
    assert ("check_mode" in payment["flags"]) is flagged


def test_a_row_with_no_amount_or_no_date_is_imported_and_flagged():
    data = build([("May 26", [row(2, "Test Theta", "9000000009", None), row(3, "Test Theta", "9000000009", 500, date=None)])])
    no_amount, no_date = data.payments
    assert no_amount["state"] == "unknown" and "unknown_state" in no_amount["flags"]
    assert no_date["state"] == "paid" and "no_date" in no_date["flags"]
    assert data.clients[0]["issue_count"] == 2


def test_one_spelling_per_service_and_online_marked():
    data = build([("May 26", [
        row(2, "A One", "9000000010", 1, description="Cons"),
        row(3, "B Two", "9000000011", 1, description="cons"),
        row(4, "C Three", "9000000012", 1, description="Cons"),
        row(5, "D Four", "9000000013", 1, description="ONLINE DIET CONS"),
    ])])
    assert [t["service"] for t in data.treatments] == ["Cons", "Cons", "Cons", "ONLINE DIET CONS"]
    assert data.treatments[1]["service_as_written"] == "cons"
    assert [t["mode"] for t in data.treatments] == ["offline", "offline", "offline", "online"]


def test_suggested_weeks_read_as_weeks():
    data = build([("May 26", [row(2, "E Five", "9000000014", 1, suggested=6), row(3, "E Five", "9000000014", 1, suggested="6w")])])
    assert [t["recommended_sessions_text"] for t in data.treatments] == ["6 weeks", "6w"]
    assert all(t["recommended_sessions"] is None for t in data.treatments)


def test_summary_says_it_is_a_revenue_sheet():
    data = build([("May 26", [row(2, "F Six", "9000000015", 1)]), ("June 26", [])])
    s = past_data.summary(data)
    assert (s["layout"], s["tabs"]) == ("revenue", ["May 26", "June 26"])


# ------------------------------------------------------------------------ the workbook

HEADER = ["S.NO", "DATE", "NAME", "MOBILE NUMBER", "LOCATION", "SOR", "SUGGESTED WEEKS", "DESCRIPTION",
          "REMARKS", "PAYMENT", "BALANCE", "NEXT PYMT DATE-MAY", "NEXT PYMT DATE-JUNE", "PAYMENT MODE",
          "PHYSIO NAME", "SLOT TIMING"]


def write_revenue_sheet(path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("Summary")  # not a month: no NAME/MOBILE/PAYMENT header
    summary.append(["Month", "Total"])
    summary.append(["May 26", 1600])
    may = wb.create_sheet("May 26")
    may.append(HEADER)
    may.append([1, datetime(2026, 5, 2), "Test Kappa", 9000000020, "Mannady", "Instagram", "6w", "Cons",
                None, 600, 0, None, None, "Gpay", None, None])
    may.append([2, datetime(2026, 5, 9), "Test Kappa", "9000000020", None, None, None, "Physio session",
                "2w", 1000, None, "16-05-2026-paid", None, "Cash", "Test Physio", "5-6PM"])
    may.append([None, None, None, None, None, None, None, None, None, 1600])  # the month's total
    june = wb.create_sheet("June 26")
    # Another branch's copy with its columns in another order, and a title row above them.
    june.append(["JUNE REVENUE"])
    june.append(["NAME", "PAYMENT", "DATE", "MOBILE NUMBER", "PAYMENT MODE", "DESCRIPTION"])
    june.append(["Test Lambda", 2000, datetime(2026, 6, 1), "9000000021", "Card", "Rehab"])
    wb.save(path)


def test_read_workbook_reads_a_revenue_sheet_and_passes_over_other_tabs(tmp_path):
    path = tmp_path / "revenue.xlsx"
    write_revenue_sheet(path)
    data = read_workbook(path)
    assert (data.layout, data.tabs) == ("revenue", ["May 26", "June 26"])
    kappa, lambda_ = by_name(data.clients)["Test Kappa"], by_name(data.clients)["Test Lambda"]
    assert kappa["treatments_count"] == 2 and kappa["paid_total"] == 1600
    assert (kappa["source"], kappa["notes"]) == ("Instagram", "Location: Mannady")
    assert lambda_["paid_total"] == 2000 and data.payments[-1]["mode"] == "card"
    physio = next(t for t in data.treatments if t["physio"])
    assert (physio["physio"], physio["notes"]) == ("Test Physio", "2w · NEXT PYMT DATE-MAY: 16-05-2026-paid · Slot 5-6PM")
    assert [f.code for f in data.findings] == ["no_name_row"]


def test_read_workbook_refuses_a_workbook_that_is_neither(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Item", "Qty", "Price"])
    path = tmp_path / "other.xlsx"
    wb.save(path)
    with pytest.raises(PastDataError, match="neither the register"):
        read_workbook(path)
