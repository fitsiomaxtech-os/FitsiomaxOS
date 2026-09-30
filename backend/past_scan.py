"""Auto Scan: what is inside a workbook, before any of it goes into Past Data.

Add Sheet reads an upload in two passes. This is the first: every tab, the row its headers
sit on, and each column under them -- the header as written, a few of its values, how many
rows fill it -- so the screen can show the whole file and have each tab and each column
turned on or off. The second pass, Fetch, reads the file through read_picked(): only the tabs
and the columns left on.

What turning a column off does depends on the file. The three sheets the OS knows by their
layout -- the register (past_data.py), a branch's revenue sheet (past_revenue.py), the OS
Data workbook (past_os.py) -- are read by their own readers, which look for their own
headers. A column turned off is emptied before they look, so it reads as a column the file
never had. The columns a reader cannot do without (the register's Patient ID, a revenue tab's
MOBILE) come back `required` and cannot be turned off; a column the reader has no use for
comes back `used: False`, since on or off, nothing would read it.

Any other workbook is a "custom" sheet -- a list of people, one per row, read by
past_custom.py. There every column is used: each is given an OS field (Name, Phone, City...),
matched here off its header, or kept as an extra detail under its own header, and the screen
can change either.

Nothing here touches a database -- see past_data.py.
"""
import re
from typing import Any, Dict, List, Optional, Tuple

import past_data
from past_data import HEADER_SEARCH_ROWS, PastDataError, amount, phone, squash, text

CUSTOM = "custom"
LAYOUT_NAMES = {"register": "register", "revenue": "revenue sheet", "os": "OS Data workbook", CUSTOM: "sheet"}

# How far down a tab of no known layout its header row is looked for: a title and a line or
# two of notes above the table is common; ten rows of it is not.
SCAN_ROWS = 10
SAMPLES = 3
SAMPLE_CHARS = 40
# How many of a column's values are looked at to tell a column of phone numbers or emails
# whose header says nothing the OS knows.
KIND_LOOK = 50

# The OS fields a custom sheet's column can go into -- the lead's own (V3LeadCreate), plus
# Date of Birth (the lead's Age is worked out from it), Status and Current Stage (where Move
# to live puts the client) -- and the headers each is known by, squash()ed. Exact matches
# only: "Company name" holds the word "name" and is not the client's name.
FIELDS: Tuple[Tuple[str, str, Tuple[str, ...]], ...] = (
    ("name", "Name", ("name", "fullname", "patientname", "clientname", "customername", "leadname",
                      "membername", "personname")),
    ("phone", "Phone", ("phone", "phoneno", "phonenumber", "mobile", "mobileno", "mobilenumber",
                        "contact", "contactno", "contactnumber", "whatsapp", "whatsappno",
                        "whatsappnumber", "cell", "cellno")),
    ("alternative_phone", "Alternative Phone", ("alternativephone", "alternatephone", "altphone",
                                                "alternativenumber", "alternatenumber", "alternativemobile",
                                                "alternatemobile", "secondaryphone", "phone2", "mobile2",
                                                "landline")),
    ("email", "Email", ("email", "emailid", "emailaddress", "mail", "mailid")),
    ("gender", "Gender", ("gender", "sex")),
    ("age", "Age", ("age",)),
    ("dob", "Date of Birth", ("dob", "dateofbirth", "birthdate", "birthday")),
    ("address", "Address", ("address", "fulladdress", "residentialaddress", "streetaddress")),
    ("location", "Location", ("location", "area", "locality")),
    ("city", "City", ("city", "town")),
    ("state", "State", ("state",)),
    ("occupation", "Occupation", ("occupation", "profession", "job", "jobtitle", "designation")),
    ("department", "Department", ("department", "dept")),
    ("condition", "Condition / Pain Area", ("condition", "conditionpainarea", "painarea", "complaint",
                                            "chiefcomplaint", "problem", "diagnosis")),
    ("months_of_pain", "Months of Pain", ("monthsofpain", "painduration")),
    ("source", "Source", ("source", "leadsource", "sor", "referredby", "reference", "referral")),
    ("registration_date", "Lead Date", ("date", "leaddate", "leaddatetime", "registrationdate", "regdate",
                                        "joiningdate", "joindate", "dateofjoining", "enquirydate",
                                        "inquirydate", "createdon", "createddate")),
    ("status", "Status", ("status", "activeorinactive", "activeinactive", "clientstatus", "active")),
    ("current_stage", "Current Stage", ("currentstage", "stage", "leadstage")),
    ("notes", "Notes", ("notes", "note", "remarks", "remark", "comments", "comment")),
)
FIELD_LABELS = {key: label for key, label, _ in FIELDS}
_BY_HEADER = {alias: key for key, _, aliases in FIELDS for alias in aliases}
# A custom tab is a list of people only with one of these given to a column.
IDENTITY_FIELDS = ("name", "phone")
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


# ------------------------------------------------------------------------ the workbook, held

class GridSheet:
    """One tab held as its rows, standing in for an openpyxl worksheet: what the readers ask
    of one is its title and iter_rows(values_only=True)."""

    def __init__(self, title: str, rows: List[tuple], hidden: bool = False):
        self.title = title
        self.rows = rows
        self.hidden = hidden

    def iter_rows(self, min_row: int = 1, max_row: Optional[int] = None, values_only: bool = True):
        yield from self.rows[min_row - 1:max_row]


class GridBook:
    """A workbook held in memory, answering what the readers ask of an openpyxl one."""

    def __init__(self, sheets: List[GridSheet]):
        self.worksheets = sheets

    @property
    def sheetnames(self) -> List[str]:
        return [ws.title for ws in self.worksheets]

    def __getitem__(self, name: str) -> GridSheet:
        for ws in self.worksheets:
            if ws.title == name:
                return ws
        raise KeyError(name)

    def close(self) -> None:
        pass


def _filled(value: Any) -> bool:
    return text(value) != ""


def grid(workbook) -> GridBook:
    """Every tab of an open workbook, read once into memory. The empty rows at the foot of a
    tab -- a template formatted down to row 1000 -- are left behind."""
    sheets = []
    for ws in workbook.worksheets:
        rows = [tuple(r or ()) for r in ws.iter_rows(values_only=True)]
        while rows and not any(_filled(c) for c in rows[-1]):
            rows.pop()
        sheets.append(GridSheet(ws.title, rows, getattr(ws, "sheet_state", "visible") != "visible"))
    return GridBook(sheets)


def _letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def _cell(row: tuple, index: int) -> Any:
    return row[index] if index < len(row) else None


def column_label(header_row: tuple, index: int) -> str:
    """A column as the screen and the lead popup name it: its header, or its letter when it
    has none."""
    return text(_cell(header_row, index)) or f"Column {_letter(index)}"


# -------------------------------------------------------------------------- which layout

def layout_of(book) -> str:
    """"register" | "os" | "revenue" | "custom", asked in the order past_data.read_open asks."""
    if past_data.CLIENT_SHEET in book.sheetnames:
        return "register"
    import past_os
    if past_os.detect(book):
        return "os"
    import past_revenue
    if any(past_revenue._header(list(ws.iter_rows(values_only=True))) for ws in book.worksheets):
        return "revenue"
    return CUSTOM


def header_row(rows: List[tuple]) -> Optional[int]:
    """The row a tab of no known layout keeps its headers on: of its first rows, the one with
    the most cells that read as a label -- words, not a figure or a date -- a header the OS
    knows (Name, Phone...) counting three times over, so a row of data full of names and
    towns does not outscore the headers above it. None for a tab with no such row."""
    best, best_score = None, 0
    for at, row in enumerate(rows[:SCAN_ROWS]):
        labels = [c for c in row if isinstance(c, str) and c.strip() and amount(c) is None]
        score = len(labels) + sum(2 for c in labels if squash(c) in _BY_HEADER)
        if score > best_score:
            best, best_score = at, score
    return best


def _reading(book, layout: str) -> Dict[str, dict]:
    """Tab title -> how `layout`'s reader reads that tab: its header row's index ("at"), the
    columns it reads ("used", None for every column), the ones it cannot do without
    ("required"), and whether the tab itself is one it cannot do without. A tab the reader
    passes over is not in it."""
    out: Dict[str, dict] = {}
    if layout == "register":
        sheets = {
            past_data.CLIENT_SHEET: past_data.CLIENT_COLUMNS,
            past_data.TREATMENT_SHEET: past_data.TREATMENT_COLUMNS,
            past_data.PAYMENT_SHEET: past_data.PAYMENT_COLUMNS,
            past_data.ENQUIRY_SHEET: past_data.ENQUIRY_COLUMNS,
        }
        for title, columns in sheets.items():
            if title not in book.sheetnames:
                continue
            rows = book[title].rows
            id_header = squash(next(iter(columns.values())))
            at = next((i for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]) if id_header in {squash(c) for c in row}), None)
            if at is None:
                continue
            position: Dict[str, int] = {}
            for i, cell in enumerate(rows[at]):
                position.setdefault(squash(cell), i)
            where = {key: position.get(squash(header)) for key, header in columns.items()}
            out[title] = {
                "at": at,
                "used": {i for i in where.values() if i is not None},
                "required": {where[k] for k in past_data.REQUIRED[title] if where[k] is not None},
                "required_tab": True,
            }
    elif layout == "os":
        import past_os
        for name in past_os.TAB_ORDER:
            ws = past_os._tab(book, name)
            found = past_os._header(ws.rows, past_os.COLUMNS[name]) if ws is not None else None
            if not found:
                continue
            at, where = found
            need = list(past_os.REQUIRED[name])
            if name == past_os.LEADS:
                # What tells the workbook for the OS Data one (past_os.detect), and the ID the
                # other tabs find their client by.
                need += ["excel_id", "current_stage"]
            else:
                need.append("client_excel_id" if where.get("client_excel_id") is not None else "phone")
            out[ws.title] = {
                "at": at,
                "used": {i for i in where.values() if i is not None},
                "required": {where[k] for k in need if where.get(k) is not None},
                "required_tab": name == past_os.LEADS,
            }
    elif layout == "revenue":
        import past_revenue
        for ws in book.worksheets:
            found = past_revenue._header(ws.rows)
            if not found:
                continue
            at, where, nexts = found
            out[ws.title] = {
                "at": at,
                "used": {i for i in where.values() if i is not None} | {i for i, _ in nexts},
                "required": {where[k] for k in past_revenue.REQUIRED},
                "required_tab": False,
            }
    else:
        for ws in book.worksheets:
            at = header_row(ws.rows)
            if at is not None:
                out[ws.title] = {"at": at, "used": None, "required": set(), "required_tab": False}
    return out


# ------------------------------------------------------------------------------ the scan

def _columns(rows: List[tuple], at: int) -> List[Tuple[int, str, List[Any]]]:
    """(index, header, the filled values under it) of each column holding a header or data."""
    header, body = rows[at], rows[at + 1:]
    width = max((len(r) for r in rows[at:]), default=0)
    found = []
    for i in range(width):
        head = text(_cell(header, i))
        values = [r[i] for r in body if i < len(r) and _filled(r[i])]
        if head or values:
            found.append((i, head, values))
    return found


def _samples(values: List[Any]) -> List[str]:
    seen: List[str] = []
    for value in values:
        shown = text(value)[:SAMPLE_CHARS]
        if shown not in seen:
            seen.append(shown)
        if len(seen) == SAMPLES:
            break
    return seen


def _looks_phone(value: Any) -> bool:
    return bool(phone(value)[1])


def _looks_email(value: Any) -> bool:
    return bool(EMAIL.match(text(value)))


def suggest(columns: List[Tuple[int, str, List[Any]]]) -> Dict[int, str]:
    """Column index -> the OS field its header names, each field given once, to the first
    column that names it. A column of phone numbers or emails under a header the OS does not
    know ("Contact 1", "WhatsApp No.") gets Phone or Email off its values, when no header took
    it first."""
    out: Dict[int, str] = {}
    for i, head, _ in columns:
        key = _BY_HEADER.get(squash(head))
        if key and key not in out.values():
            out[i] = key
    for key, looks in (("phone", _looks_phone), ("email", _looks_email)):
        if key in out.values():
            continue
        for i, head, values in columns:
            sample = values[:KIND_LOOK]
            if i not in out and head and sample and sum(1 for v in sample if looks(v)) * 10 >= len(sample) * 6:
                out[i] = key
                break
    return out


def scan(book) -> Dict[str, Any]:
    """What the workbook holds, tab by tab and column by column, and what each is set to
    before anybody changes it: every tab the reader reads, on (a hidden one or an empty one
    off, unless the reader cannot do without it); every column the reader reads, on; and, on
    a custom sheet, each column's suggested OS field."""
    layout = layout_of(book)
    reading = _reading(book, layout)
    tabs = []
    for ws in book.worksheets:
        how = reading.get(ws.title)
        entry: Dict[str, Any] = {
            "name": ws.title, "hidden": ws.hidden, "header_row": None, "rows": 0,
            "used": how is not None, "required": bool(how and how["required_tab"]),
            "on": False, "note": "", "columns": [],
        }
        if not how:
            entry["note"] = "Empty" if not ws.rows else (
                "No header row found" if layout == CUSTOM else f"Not part of the {LAYOUT_NAMES[layout]}")
            tabs.append(entry)
            continue
        at = how["at"]
        columns = _columns(ws.rows, at)
        used = how["used"]
        fields = suggest(columns) if layout == CUSTOM else {}
        counted = {i for i, _, _ in columns} if used is None else used
        entry["header_row"] = at + 1
        entry["rows"] = sum(1 for r in ws.rows[at + 1:] if any(_filled(_cell(r, i)) for i in counted))
        entry["on"] = entry["required"] or (entry["rows"] > 0 and not ws.hidden)
        for i, head, values in columns:
            reads = used is None or i in used
            entry["columns"].append({
                "index": i,
                "letter": _letter(i),
                "header": head,
                "samples": _samples(values),
                "filled": len(values),
                "used": reads,
                "required": i in how["required"],
                # A custom sheet's column with no header is somebody's scribbles more often
                # than data; it is listed, and left off.
                "on": reads and (used is not None or bool(head)),
                "field": fields.get(i, ""),
            })
        tabs.append(entry)
    return {
        "layout": layout,
        "tabs": tabs,
        "fields": [{"key": key, "label": label} for key, label, _ in FIELDS] if layout == CUSTOM else [],
    }


def suggested_picks(book) -> Dict[str, dict]:
    """A custom sheet read as Auto Scan first shows it: every tab it turns on that has a Name
    or Phone column, every headed column, each under its suggested field -- how the terminal
    tool reads one."""
    picks = {}
    for tab in scan(book)["tabs"]:
        fields = {c["index"]: c["field"] for c in tab["columns"] if c["on"]}
        if tab["on"] and set(fields.values()) & set(IDENTITY_FIELDS):
            picks[tab["name"]] = {"at": tab["header_row"] - 1, "fields": fields}
    return picks


# ----------------------------------------------------------------------------- the fetch

def _checked(book, picks: Any, layout: str, reading: Dict[str, dict]) -> Dict[str, Dict[int, str]]:
    """The picks as {tab: {column index: field}}, once they hold together: every tab and
    column one the scan listed, nothing the reader needs left off, and on a custom sheet each
    field given once a tab and every tab left on naming somebody. PastDataError otherwise,
    saying which."""
    if not isinstance(picks, list):
        raise PastDataError("the columns picked could not be read -- scan the file again")
    chosen: Dict[str, Dict[int, str]] = {}
    for pick in picks:
        title = pick.get("tab") if isinstance(pick, dict) else None
        if title not in book.sheetnames:
            raise PastDataError(f"there is no '{title}' tab in this file")
        if title not in reading:
            raise PastDataError(f"'{title}' has no table to read -- turn it off")
        if title in chosen:
            raise PastDataError(f"'{title}' is picked twice")
        listed = {i for i, _, _ in _columns(book[title].rows, reading[title]["at"])}
        fields: Dict[int, str] = {}
        for column in pick.get("columns") or []:
            index = column.get("index") if isinstance(column, dict) else None
            if not isinstance(index, int) or isinstance(index, bool) or index not in listed:
                raise PastDataError(f"'{title}' has no column {index}")
            field = (column.get("field") or "") if layout == CUSTOM else ""
            if field and field not in FIELD_LABELS:
                raise PastDataError(f"'{field}' is not an OS field")
            if field and field in fields.values():
                raise PastDataError(f"'{title}': {FIELD_LABELS[field]} is given to two columns -- keep it on one")
            fields[index] = field
        chosen[title] = fields
    if not chosen:
        raise PastDataError("every tab is turned off -- turn on the ones to fetch")
    for title, how in reading.items():
        if how["required_tab"] and title not in chosen:
            raise PastDataError(f"'{title}' is needed to read the {LAYOUT_NAMES[layout]} -- turn it back on")
        header = book[title].rows[how["at"]]
        off = [column_label(header, i) for i in sorted(how["required"]) if title in chosen and i not in chosen[title]]
        if off:
            raise PastDataError(f"'{title}': {', '.join(off)} {'is' if len(off) == 1 else 'are'} needed to read it -- turn "
                                f"{'it' if len(off) == 1 else 'them'} back on")
    if layout == CUSTOM:
        for title, fields in chosen.items():
            if not set(fields.values()) & set(IDENTITY_FIELDS):
                raise PastDataError(f"'{title}': give one column Name or Phone, or turn the tab off")
    return chosen


def _blanked(book, chosen: Dict[str, Dict[int, str]], reading: Dict[str, dict]) -> GridBook:
    """The workbook as a known layout's reader should see it: only the tabs left on, and in
    each, the columns turned off emptied from the header row down. Above the header (a title,
    a row of totals) is left as it is, so the header is found where it always was."""
    sheets = []
    for ws in book.worksheets:
        if ws.title not in chosen:
            continue
        keep, at = chosen[ws.title], reading[ws.title]["at"]
        rows = ws.rows[:at] + [tuple(c if i in keep else None for i, c in enumerate(r)) for r in ws.rows[at:]]
        sheets.append(GridSheet(ws.title, rows, ws.hidden))
    return GridBook(sheets)


def read_picked(book: GridBook, picks: Any) -> past_data.PastData:
    """Read a workbook through the tabs and columns picked on its scan. `picks` is
    [{"tab": title, "columns": [{"index": i, "field": key or ""}]}] -- the tabs left on and,
    in each, the columns left on; `field` is read on a custom sheet only."""
    layout = layout_of(book)
    reading = _reading(book, layout)
    chosen = _checked(book, picks, layout, reading)
    if layout == CUSTOM:
        import past_custom
        data = past_custom.read(book, {t: {"at": reading[t]["at"], "fields": f} for t, f in chosen.items()})
    else:
        data = past_data.read_open(_blanked(book, chosen, reading), layout=layout)
    data.columns = [
        {"tab": title, "columns": [
            {"header": column_label(book[title].rows[reading[title]["at"]], i), "field": field}
            for i, field in sorted(fields.items())
        ]}
        for title, fields in chosen.items()
    ]
    return data


def scan_file(source) -> Dict[str, Any]:
    """scan() of a path or an open binary file."""
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
    try:
        return scan(grid(workbook))
    finally:
        workbook.close()
